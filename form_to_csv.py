from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from datetime import datetime, date


def get_finished_parcels(schlag_collection):
    finished_parcels = []

    for s_id, s in schlag_collection.items():
        if "Düngeinformation" in s.keys() and s["Düngeinformation"]:
            finished_parcels.append(s_id)

    return finished_parcels


def build_dbe_excel(schlag_collection, metadata):
    """Create a formatted English nitrogen requirement report."""

    wb = Workbook()
    summary = wb.active
    summary.title = "Overview"

    green = "245640"
    light_green = "E2F0E8"
    pale = "F5F8F5"
    white = "FFFFFF"
    grey = "667085"
    border_color = "D0D5DD"

    thin = Side(style="thin", color=border_color)

    def setup_sheet(ws):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A5"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_margins = PageMargins(
            left=0.25,
            right=0.25,
            top=0.5,
            bottom=0.5,
        )
        ws.oddFooter.center.text = "Nitrogen Requirement Report"
        ws.oddFooter.right.text = "Page &P of &N"

    def title_block(ws, title, subtitle, last_col):
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=last_col,
        )

        ws["A1"] = title
        ws["A1"].font = Font(
            name="Aptos Display",
            size=18,
            bold=True,
            color=white,
        )
        ws["A1"].fill = PatternFill("solid", fgColor=green)
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 36

        for col in range(1, last_col + 1):
            ws.cell(1, col).fill = PatternFill(
                "solid",
                fgColor=green,
            )

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=last_col,
        )

        ws["A2"] = subtitle
        ws["A2"].font = Font(
            size=10,
            italic=True,
            color=grey,
        )
        ws["A2"].alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )
        ws.row_dimensions[2].height = 24

    def write_table(ws, headers, rows, table_name):
        header_row = 4

        for col, header in enumerate(headers, 1):
            cell = ws.cell(header_row, col, header)
            cell.fill = PatternFill("solid", fgColor=green)
            cell.font = Font(bold=True, color=white)
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
            )
            cell.border = Border(bottom=thin)

        ws.row_dimensions[header_row].height = 42

        for row_idx, values in enumerate(rows, 5):
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row_idx, col_idx, value)
                cell.fill = PatternFill(
                    "solid",
                    fgColor=white if row_idx % 2 else pale,
                )
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

                if isinstance(value, (int, float)):
                    cell.number_format = "#,##0.00"

                elif isinstance(value, (date, datetime)):
                    cell.number_format = "DD.MM.YYYY"

        last_col = get_column_letter(len(headers))
        last_row = max(4, 4 + len(rows))

        if rows:
            tab = Table(
                displayName=table_name,
                ref=f"A4:{last_col}{last_row}",
            )

            tab.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showRowStripes=True,
                showColumnStripes=False,
            )

            ws.add_table(tab)

        ws.freeze_panes = "A5"
        ws.print_title_rows = "1:4"

        for col_idx, header in enumerate(headers, 1):
            width = min(max(len(str(header)) + 3, 15), 32)
            ws.column_dimensions[
                get_column_letter(col_idx)
            ].width = width

    # --------------------------------------------------
    # 1. Collect saved parcel and calculation data
    # --------------------------------------------------

    parcel_rows = []
    detail_rows = []
    fertilizer_rows = []

    for s_id, s in schlag_collection.items():

        area = s.get("Größe in ha")
        area = float(area) if area is not None else None

        dbe_result = s.get(
            "Stickstoffbedarf durch Dünger in kg / ha"
        )

        args = s.get("DBE_Args", {})
        main_crop = args.get("Hauptfrucht", {})
        catch_crop = args.get("Zwischenfrucht", {})

        n_total_ha = 0.0
        n_total_parcel = 0.0

        fertilizer_data = s.get("Düngeinformation", {}) or {}

        for d_id, f in fertilizer_data.items():
            if not isinstance(f, dict):
                continue

            amount_key = next(
                (
                    k for k in f
                    if k.startswith("Düngemente in ")
                ),
                None,
            )

            total_key = next(
                (
                    k for k in f
                    if k.startswith("Stickstoff Gesamt in ")
                ),
                None,
            )

            available_key = next(
                (
                    k for k in f
                    if k.startswith(
                        "Stickstoff Pflanzenverfügbar in "
                    )
                ),
                None,
            )

            total_ha = f.get(total_key) if total_key else None

            available_ha = (
                f.get(available_key)
                if available_key
                else None
            )

            if isinstance(total_ha, (int, float)):
                n_total_ha += total_ha

                if area is not None:
                    n_total_parcel += total_ha * area

            fertilizer_rows.append([
                s.get("Schlag"),
                s.get("Feldblock-Nummer"),
                area,
                d_id,
                f.get("Datum"),
                f.get("Düngetyp"),
                f.get("Einheit"),
                f.get(amount_key) if amount_key else None,
                total_ha,
                (
                    total_ha * area
                    if isinstance(total_ha, (int, float))
                    and area is not None
                    else None
                ),
                available_ha,
                (
                    available_ha * area
                    if isinstance(available_ha, (int, float))
                    and area is not None
                    else None
                ),
            ])

        # One row per parcel in the main calculation sheet.
        parcel_rows.append([
            s_id,
            s.get("Schlag"),
            s.get("Feldblock-Nummer"),
            area,
            s.get("Kultur"),
            s.get("Vorjahreskultur"),
            (
                "None"
                if s.get("Zwischenfrucht") == "Keine"
                else s.get("Zwischenfrucht")
            ),
            s.get("Datum der Aussaat"),
            (
                "None"
                if s.get("Datum der Aussaat Zwischenfrucht") == "Keine"
                else s.get("Datum der Aussaat Zwischenfrucht")
            ),
            main_crop.get("Ertrag laut DüV"),
            main_crop.get("Durchschnittsertrag 5 Jahre Betrieb"),
            main_crop.get("N-Bedarfswert"),
            main_crop.get("Abschlag Ertragsdifferenz"),
            main_crop.get("N_min Probe / Richtwert"),
            main_crop.get("Standort / Humus"),
            main_crop.get("Abschlag Vorfrucht"),
            main_crop.get("Abschlag Zwischenfrucht"),
            dbe_result,
            n_total_ha,
            n_total_parcel if area is not None else None,
            n_total_ha / 10,
            (
                "Calculated"
                if dbe_result is not None
                else "Not yet calculated"
            ),
        ])

        # Audit trail: one row per saved calculation parameter.
        for section, parameters in args.items():
            if not isinstance(parameters, dict):
                continue

            for parameter, value in parameters.items():
                translated_section = {
                    "Hauptfrucht": "Main Crop",
                    "Zwischenfrucht": "Catch Crop",
                }.get(section, section)

                translated_parameter = {
                    "Datum": "Sowing Date",
                    "Kultur": "Crop",
                    "Ertrag laut DüV":
                        "Yield Reference under German Fertilizer Ordinance",
                    "Durchschnittsertrag 5 Jahre Betrieb":
                        "Farm Average Yield over 5 Years",
                    "N-Bedarfswert":
                        "Nitrogen Requirement Reference Value",
                    "Abschlag Ertragsdifferenz":
                        "Yield Difference Adjustment",
                    "N_min Probe / Richtwert":
                        "Nmin Measurement / Reference Value",
                    "Standort / Humus":
                        "Soil / Humus Adjustment",
                    "Abschlag Vorfrucht":
                        "Previous Crop Deduction",
                    "Abschlag Zwischenfrucht":
                        "Catch Crop Deduction",
                    "N-Bedarf":
                        "Nitrogen Requirement",
                }.get(parameter, parameter)

                detail_rows.append([
                    s.get("Schlag"),
                    s.get("Feldblock-Nummer"),
                    translated_section,
                    translated_parameter,
                    (
                        value.isoformat()
                        if isinstance(value, (date, datetime))
                        else (
                            "None"
                            if value == "Keine"
                            else value
                        )
                    ),
                ])

    # --------------------------------------------------
    # 2. Overview sheet
    # --------------------------------------------------

    setup_sheet(summary)

    title_block(
        summary,
        "Nitrogen Fertilizer Requirement",
        "Report and supporting records based on saved calculation data",
        4,
    )

    summary["A4"] = "REPORT INFORMATION"
    summary["A4"].font = Font(
        bold=True,
        color=green,
        size=12,
    )

    for row, (label, value) in enumerate(metadata, 5):
        summary.cell(row, 1, label).font = Font(bold=True)

        summary.cell(row, 1).fill = PatternFill(
            "solid",
            fgColor=light_green,
        )

        if isinstance(value, datetime):
            value = value.strftime("%d.%m.%Y %H:%M")
        elif isinstance(value, date):
            value = value.strftime("%d.%m.%Y")

        summary.cell(row, 2, value)

    completed = sum(
        1
        for s in schlag_collection.values()
        if s.get("Stickstoffbedarf durch Dünger in kg / ha")
        is not None
    )

    total_area = sum(
        float(s["Größe in ha"])
        for s in schlag_collection.values()
        if isinstance(s.get("Größe in ha"), (int, float))
    )

    summary["A11"] = "KEY FIGURES"
    summary["A11"].font = Font(
        bold=True,
        color=green,
        size=12,
    )

    metrics = [
        ("Registered Parcels", len(schlag_collection)),
        ("Parcels with Calculated Requirements", completed),
        ("Total Parcel Area (ha)", total_area),
        ("Documented Fertilization Events", len(fertilizer_rows)),
    ]

    for row, (label, value) in enumerate(metrics, 12):
        summary.cell(row, 1, label)
        summary.cell(row, 2, value)

        summary.cell(row, 1).font = Font(bold=True)

        summary.cell(row, 1).fill = PatternFill(
            "solid",
            fgColor=light_green,
        )

        summary.cell(row, 2).number_format = "#,##0.00"

    summary["A18"] = "NOTES"
    summary["A18"].font = Font(
        bold=True,
        color=green,
        size=12,
    )

    notes = [
        "Nitrogen requirement results are imported from the application.",
        "Blank values do not automatically mean zero.",
        "Nitrogen requirement values are expressed per hectare; "
        "parcel totals are reported separately.",
        "This workbook documents the saved data. "
        "Legal compliance and completeness must be checked separately.",
    ]

    for row, note in enumerate(notes, 19):
        summary.cell(row, 1, f"• {note}")

        summary.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=4,
        )

        summary.cell(row, 1).alignment = Alignment(wrap_text=True)
        summary.row_dimensions[row].height = 28

    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 30
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 18

    # --------------------------------------------------
    # 3. Nitrogen requirement by parcel
    # --------------------------------------------------

    dbe = wb.create_sheet("Requirements by Parcel")
    setup_sheet(dbe)

    title_block(
        dbe,
        "Nitrogen Requirement by Parcel",
        "Calculation inputs, adjustments and calculated nitrogen requirements",
        22,
    )

    dbe_headers = [
        "Parcel ID",
        "Parcel Name",
        "Field Block Number",
        "Area (ha)",
        "Current Crop",
        "Previous Year's Crop",
        "Catch Crop",
        "Main Crop Sowing Date",
        "Catch Crop Sowing Date",
        "Yield Reference under German Fertilizer Ordinance (dt/ha)",
        "5-Year Average Yield (dt/ha)",
        "Nitrogen Requirement Reference Value (kg N/ha)",
        "Yield Difference Adjustment (kg N/ha)",
        "Nmin (kg N/ha)",
        "Humus Adjustment (kg N/ha)",
        "Previous Crop Deduction (kg N/ha)",
        "Catch Crop Deduction (kg N/ha)",
        "Nitrogen Fertilizer Requirement (kg N/ha)",
        "Total Nitrogen Applied (kg N/ha)",
        "Total Nitrogen Applied (kg N/parcel)",
        "Following-Year Deduction (kg N/ha)",
        "Status",
    ]

    write_table(
        dbe,
        dbe_headers,
        parcel_rows,
        "DBEParcelTable",
    )

    # --------------------------------------------------
    # 4. Fertilization records
    # --------------------------------------------------

    fert = wb.create_sheet("Fertilization Records")
    setup_sheet(fert)

    title_block(
        fert,
        "Nitrogen Fertilization Records",
        "Recorded fertilization events and nitrogen quantities",
        12,
    )

    fert_headers = [
        "Parcel Name",
        "Field Block Number",
        "Area (ha)",
        "Fertilization ID",
        "Date",
        "Fertilizer",
        "Unit",
        "Fertilizer Amount per ha",
        "Total Nitrogen (kg N/ha)",
        "Total Nitrogen (kg N/parcel)",
        "Plant-Available Nitrogen (kg N/ha)",
        "Plant-Available Nitrogen (kg N/parcel)",
    ]

    write_table(
        fert,
        fert_headers,
        fertilizer_rows,
        "FertilizerTable",
    )

    # --------------------------------------------------
    # 5. Calculation details
    # --------------------------------------------------

    audit = wb.create_sheet("Calculation Details")
    setup_sheet(audit)

    title_block(
        audit,
        "Calculation Details",
        "Saved calculation inputs and parameters for each parcel",
        5,
    )

    audit_headers = [
        "Parcel Name",
        "Field Block Number",
        "Calculation Section",
        "Parameter",
        "Saved Value",
    ]

    write_table(
        audit,
        audit_headers,
        detail_rows,
        "AuditDetailTable",
    )

    # --------------------------------------------------
    # 6. Workbook formatting and export
    # --------------------------------------------------

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.tabColor = green
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.properties.title = "Nitrogen Fertilizer Requirement Report"
    wb.properties.subject = (
        "Nitrogen fertilizer requirements and supporting records"
    )
    wb.properties.creator = "Nitrogen Requirement Application"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()
